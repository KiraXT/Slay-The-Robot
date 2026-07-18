#!/usr/bin/env python3
import json
import csv
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
CARDS = ROOT / "external" / "data" / "cards"
XLSX = ROOT / "external" / "config" / "cards.xlsx"
CSV = ROOT / "external" / "config" / "cards.csv"

EXPECTED = {
    "card_hasty_sorting": "color_blue",
    "card_something_fell_out": "color_blue",
    "card_quick_search": "color_blue",
    "card_by_accident": "color_blue",
    "card_backup_drink": "color_blue",
    "card_first_beat": "color_green",
    "card_distorted_note": "color_green",
    "card_tuning": "color_green",
    "card_warmup": "color_orange",
    "card_sports_drink": "color_orange",
    "card_bag_swing": "color_orange",
    "card_old_item_reuse": "color_orange",
    "card_sprint_start": "color_orange",
    "card_opening_strike": "color_red",
    "card_finisher": "color_red",
    "card_pursuit": "color_red",
    "card_echo_shield": "color_green",
    "card_finale_burst": "color_green",
    "card_metronome": "color_green",
    "card_travel_light": "color_orange",
    "card_last_item": "color_orange",
    "card_combo_starter": "color_red",
    "card_chorus": "color_green",
    "card_bookmark_clip": "color_orange",
    "card_pack_sorting": "color_orange",
    "card_ready_stance": "color_orange",
}

REQUIRED_ACTIONS = {
    "card_something_fell_out": ["card_discard_actions"],
    "card_sports_drink": ["card_draw_actions"],
    "card_warmup": ["card_retain_actions"],
    "card_sprint_start": ["card_draw_actions"],
    "card_pursuit": ["card_play_actions"],
    "card_echo_shield": ["card_play_actions"],
    "card_finale_burst": ["card_play_actions"],
    "card_metronome": ["card_play_actions"],
    "card_travel_light": ["card_play_actions"],
    "card_last_item": ["card_play_actions"],
    "card_combo_starter": ["card_play_actions"],
    "card_chorus": ["card_play_actions"],
    "card_bookmark_clip": ["card_play_actions"],
    "card_pack_sorting": ["card_play_actions"],
    "card_ready_stance": ["card_play_actions"],
}

ALLOWED_SCRIPT_PREFIXES = (
    "res://scripts/actions/",
    "res://scripts/validators/",
    "res://scripts/card_listeners/",
)


def load_card(card_id: str) -> dict:
    path = CARDS / f"{card_id}.json"
    assert path.exists(), f"missing card JSON: {path}"
    with path.open(encoding="utf-8") as handle:
        return json.load(handle)["properties"]


def walk_scripts(value):
    if isinstance(value, dict):
        for key, child in value.items():
            if isinstance(key, str) and key.startswith("res://"):
                yield key
            yield from walk_scripts(child)
    elif isinstance(value, list):
        for child in value:
            yield from walk_scripts(child)


def workbook_card_ids() -> list[str]:
    workbook = load_workbook(XLSX, read_only=True, data_only=True)
    worksheet = workbook["Cards"]
    headers = [cell.value for cell in worksheet[1]]
    object_id_index = headers.index("object_id")
    ids = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        object_id = row[object_id_index]
        if object_id is None:
            continue
        ids.append(str(object_id))
    workbook.close()
    return ids


def csv_card_ids() -> list[str]:
    with CSV.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        return [row["object_id"] for row in reader if row.get("object_id")]


def main() -> None:
    for card_id, color_id in EXPECTED.items():
        card = load_card(card_id)
        assert card["object_id"] == card_id
        assert card["card_color_id"] == color_id
        assert card["card_appears_in_card_packs"] is True
        assert card["card_rarity"] in (1, 2, 3)
        assert card["card_type"] in (0, 1, 2)
        assert card["card_texture_path"], f"{card_id} should have a texture path"
        for script_path in walk_scripts(card):
            assert script_path.startswith(ALLOWED_SCRIPT_PREFIXES), script_path
            local_path = ROOT / script_path.replace("res://", "")
            assert local_path.exists(), f"{card_id} references missing script: {script_path}"

    for card_id, fields in REQUIRED_ACTIONS.items():
        card = load_card(card_id)
        for field in fields:
            assert field in card, f"{card_id} missing {field}"
            assert len(card[field]) > 0, f"{card_id} has empty {field}"

    xlsx_ids = workbook_card_ids()
    csv_ids = csv_card_ids()
    for card_id in EXPECTED:
        assert xlsx_ids.count(card_id) == 1, f"{card_id} should appear once in cards.xlsx"
        assert csv_ids.count(card_id) == 1, f"{card_id} should appear once in cards.csv"

    print(f"validated {len(EXPECTED)} configurable cards")


if __name__ == "__main__":
    main()
