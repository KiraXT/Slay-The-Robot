#!/usr/bin/env python3
import argparse
import csv
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from tools.card_art_manifest import (
    key_color_for_card,
    manifest_card_ids,
    validate_approved_digest_contract,
    validate_key_color_contract,
)


CSV_PATH = ROOT / "external" / "config" / "cards.csv"
XLSX_PATH = ROOT / "external" / "config" / "cards.xlsx"
MANIFEST = ROOT / "tools" / "card_art_individualization_manifest.json"

EXPECTED_PHASES = {
    "pilot": {
        "card_finisher",
        "card_backup_drink",
        "card_metronome",
        "card_bag_swing",
        "card_block_initial",
        "variable_cost_attack_card",
        "card_attack_block",
        "card_damage_increase",
    },
    "red": {
        "attack_lower_cost_on_discard_card",
        "attack_with_conditional_block_card",
        "attack_with_conditional_draw_card",
        "card_banish_attack",
        "card_discard_block",
        "card_draft_red_card",
        "card_duplicate_attacks",
        "card_law",
        "card_opening_strike",
        "card_play_random_from_hand",
        "card_pursuit",
        "card_requires_adjacency",
        "card_right_click_transform_mode_a",
        "card_right_click_transform_mode_b",
        "card_vulnerable_enemies",
        "cards_played_attack_card",
        "ignore_damage_increase_attack_card",
        "set_hand_energy_card",
        "transform_hand_card",
    },
    "colored": {
        "card_by_accident",
        "card_hasty_sorting",
        "card_quick_search",
        "card_something_fell_out",
        "card_distorted_note",
        "card_echo_shield",
        "card_finale_burst",
        "card_first_beat",
        "card_tuning",
        "card_last_item",
        "card_old_item_reuse",
        "card_sports_drink",
        "card_sprint_start",
        "card_travel_light",
        "card_warmup",
    },
    "team": {
        "add_health_card",
        "card_shove",
        "custom_block_card",
        "end_turn_card",
        "attack_increase_cost_on_damage_taken_card",
        "card_energy_on_discard",
        "improving_retain_block_card",
        "upgrade_entire_deck_card",
    },
}

EXCLUDED_DEVELOPMENT_CARDS = {"card_debug_log", "card_restart_combat"}

# Current approved minima are 34,152 pixels, 268px wide, and 261px high.
# These thresholds retain at least 23% geometric and 41% area margin.
MIN_VISIBLE_PIXEL_COUNT = 20_000
MIN_VISIBLE_WIDTH = 200
MIN_VISIBLE_HEIGHT = 200
MAX_VISIBLE_WIDTH = 420
MAX_VISIBLE_HEIGHT = 420


def load_manifest() -> dict:
    assert MANIFEST.exists(), "missing card art individualization manifest"
    with MANIFEST.open(encoding="utf-8") as handle:
        return json.load(handle)


def csv_cards() -> dict[str, dict[str, str]]:
    with CSV_PATH.open(encoding="utf-8", newline="") as handle:
        return {
            row["object_id"]: row
            for row in csv.DictReader(handle)
            if row.get("object_id") and not row["object_id"].startswith("#")
        }


def xlsx_cards() -> dict[str, dict[str, str]]:
    workbook = load_workbook(XLSX_PATH, read_only=True, data_only=True)
    worksheet = workbook["Cards"]
    headers = [cell.value for cell in worksheet[1]]
    rows = {}
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        card_id = row.get("object_id")
        if card_id:
            rows[str(card_id)] = row
    workbook.close()
    return rows


def expected_art_path(card_id: str, cards: dict[str, dict[str, str]]) -> str:
    color_folder = cards[card_id]["card_color_id"].removeprefix("color_")
    return f"external/sprites/cards/{color_folder}/{card_id}.png"


def validate_card_configs(
    card_ids: list[str],
    cards: dict[str, dict[str, str]],
    workbook_cards: dict[str, dict[str, str]],
) -> None:
    for card_id in card_ids:
        assert card_id in cards, f"missing card in cards.csv: {card_id}"
        assert card_id in workbook_cards, f"missing card in cards.xlsx: {card_id}"
        assert (workbook_cards[card_id].get("card_texture_path") or "") == cards[card_id]["card_texture_path"], (
            f"card_texture_path mismatch: {card_id}"
        )
        assert expected_art_path(card_id, cards).startswith("external/sprites/cards/")


def rgb_from_hex(value: str) -> tuple[int, int, int]:
    stripped = value.removeprefix("#")
    assert len(stripped) == 6, f"invalid RGB color: {value}"
    return tuple(int(stripped[index:index + 2], 16) for index in (0, 2, 4))


def image_pixels(image: Image.Image):
    if hasattr(image, "get_flattened_data"):
        return image.get_flattened_data()
    return image.getdata()


def validate_approved_asset(
    card_id: str,
    image: Image.Image,
    approved_digest: str,
) -> None:
    assert image.size == (512, 512), (
        f"unexpected image size for {card_id}: {image.size}"
    )
    assert image.mode == "RGBA", (
        f"unexpected image mode for {card_id}: {image.mode}"
    )
    alpha_extrema = image.getextrema()[3]
    assert alpha_extrema[0] == 0, f"image is not transparent for {card_id}"
    assert alpha_extrema[1] == 255, f"image has no opaque pixels for {card_id}"

    visible_alpha = image.getchannel("A").point(
        lambda alpha: 255 if alpha > 16 else 0
    )
    visible_bbox = visible_alpha.getbbox()
    assert visible_bbox is not None, f"image has no visible subject for {card_id}"
    visible_pixel_count = visible_alpha.histogram()[255]
    assert visible_pixel_count >= MIN_VISIBLE_PIXEL_COUNT, (
        f"visible pixel count is too small for {card_id}: "
        f"{visible_pixel_count} < {MIN_VISIBLE_PIXEL_COUNT}"
    )
    visible_width = visible_bbox[2] - visible_bbox[0]
    visible_height = visible_bbox[3] - visible_bbox[1]
    assert visible_width >= MIN_VISIBLE_WIDTH, (
        f"subject is too narrow for {card_id}: {visible_bbox}"
    )
    assert visible_height >= MIN_VISIBLE_HEIGHT, (
        f"subject is too short for {card_id}: {visible_bbox}"
    )
    assert visible_width <= MAX_VISIBLE_WIDTH, (
        f"subject exceeds safe width for {card_id}: {visible_bbox}"
    )
    assert visible_height <= MAX_VISIBLE_HEIGHT, (
        f"subject exceeds safe height for {card_id}: {visible_bbox}"
    )
    for corner in ((0, 0), (511, 0), (0, 511), (511, 511)):
        assert image.getpixel(corner)[3] <= 16, (
            f"opaque corner {corner} for {card_id}"
        )

    rgba_digest = hashlib.sha256(image.convert("RGBA").tobytes()).hexdigest()
    assert rgba_digest == approved_digest, (
        f"approved RGBA digest mismatch for {card_id}: "
        f"{rgba_digest} != {approved_digest}"
    )


def validate_card_art(
    card_id: str,
    cards: dict[str, dict[str, str]],
    workbook_cards: dict[str, dict[str, str]],
    manifest: dict,
) -> None:
    expected_path = expected_art_path(card_id, cards)
    assert cards[card_id]["card_texture_path"] == expected_path, (
        f"unexpected cards.csv art path for {card_id}: "
        f"{cards[card_id]['card_texture_path']}"
    )

    assert (workbook_cards[card_id].get("card_texture_path") or "") == expected_path, (
        f"unexpected cards.xlsx art path for {card_id}: "
        f"{workbook_cards[card_id].get('card_texture_path')}"
    )

    image_path = ROOT / expected_path
    assert image_path.exists(), f"missing individual card art: {expected_path}"
    with Image.open(image_path) as image:
        validate_approved_asset(
            card_id,
            image,
            manifest["approved_rgba_sha256"][card_id],
        )

        key_hex = key_color_for_card(
            manifest,
            card_id,
            cards[card_id]["card_color_id"],
        )
        key_rgb = rgb_from_hex(key_hex)
        fringe_pixels = 0
        for red, green, blue, alpha in image_pixels(image):
            if 0 < alpha < 255 and max(
                abs(red - key_rgb[0]),
                abs(green - key_rgb[1]),
                abs(blue - key_rgb[2]),
            ) <= 32:
                fringe_pixels += 1
        assert fringe_pixels < 64, (
            f"chroma fringe detected for {card_id}: "
            f"{fringe_pixels} pixels near {key_hex}"
        )


def validate_manifest_invariants(manifest: dict) -> list[str]:
    assert sum(map(len, EXPECTED_PHASES.values())) == 50
    assert not set.union(*EXPECTED_PHASES.values()) & EXCLUDED_DEVELOPMENT_CARDS

    assert set(manifest["phases"]) == set(EXPECTED_PHASES)
    assert set(manifest["excluded_development_cards"]) == EXCLUDED_DEVELOPMENT_CARDS

    phase_ids = manifest_card_ids(manifest)
    assert len(phase_ids) == 50
    assert len(set(phase_ids)) == 50
    assert not set(phase_ids) & EXCLUDED_DEVELOPMENT_CARDS, (
        "manifest phases include excluded development cards"
    )
    for phase_name, expected_ids in EXPECTED_PHASES.items():
        assert set(manifest["phases"][phase_name]) == expected_ids
    validate_key_color_contract(manifest, phase_ids)
    validate_approved_digest_contract(manifest, phase_ids)
    return phase_ids


def validate_manifest() -> None:
    phase_ids = validate_manifest_invariants(load_manifest())

    cards = csv_cards()
    validate_card_configs(phase_ids, cards, xlsx_cards())


def validate_cards(card_ids: list[str], manifest: dict) -> None:
    cards = csv_cards()
    workbook_cards = xlsx_cards()
    validate_card_configs(card_ids, cards, workbook_cards)
    for card_id in card_ids:
        validate_card_art(card_id, cards, workbook_cards, manifest)


def validate_phase(phase_name: str) -> None:
    assert phase_name in EXPECTED_PHASES, f"unknown phase: {phase_name}"
    manifest = load_manifest()
    validate_manifest_invariants(manifest)
    phase_ids = manifest["phases"][phase_name]
    assert set(phase_ids) == EXPECTED_PHASES[phase_name]
    validate_cards(phase_ids, manifest)


def validate_all() -> None:
    manifest = load_manifest()
    phase_ids = validate_manifest_invariants(manifest)
    validate_cards(phase_ids, manifest)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--phase", choices=("manifest", "all", *EXPECTED_PHASES))
    args = parser.parse_args()
    if args.phase == "manifest":
        validate_manifest()
        print("MANIFEST_VALIDATED: 50 cards")
    elif args.phase == "all":
        validate_all()
        print("CARD_ART_VALIDATED: 50 cards")
    elif args.phase:
        validate_phase(args.phase)
        print(f"PHASE_VALIDATED: {args.phase}")
    else:
        validate_manifest()
        print("MANIFEST_VALIDATED: 50 cards")


if __name__ == "__main__":
    main()
